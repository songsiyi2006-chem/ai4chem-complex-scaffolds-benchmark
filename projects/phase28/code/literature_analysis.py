"""Descriptive reanalysis of public numerical figure-source data; no rate fitting."""
from pathlib import Path
import csv
import gzip
import hashlib
import json
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from analysis import write_csv,dump


def extract_workbooks(source: Path, destination: Path):
    """Called explicitly; preserves every valid pair and Excel row provenance.

    Unit/caption audit lives in trace_metadata.json. The source workbooks are fetched
    by reproducible URL + hash, rather than bundled as third-party article/SI copies.
    """
    import openpyxl
    destination.mkdir(parents=True,exist_ok=True)
    mappings=[("nature2023_41586_2023_5886_MOESM2_ESM.xlsx","Figure2g",0,1,"nature2023_constant_potential"),
              ("nature2023_41586_2023_5886_MOESM2_ESM.xlsx","Figure2g",3,4,"nature2023_constant_current"),
              ("natsynth2026_44160_2026_1039_MOESM4_ESM.xlsx","2h",0,1,"natsynth2026_constant_current")]
    manifests=[]
    for filename,sheetname,colx,coly,name in mappings:
        path=source/filename
        book=openpyxl.load_workbook(path,data_only=True,read_only=True)
        rows=[]
        for excel_row,row in enumerate(book[sheetname].iter_rows(values_only=True),start=1):
            x,y=row[colx],row[coly]
            if isinstance(x,(float,int)) and isinstance(y,(float,int)):
                rows.append({"excel_row":excel_row,"x":x,"y":y})
        target=destination/(name+".csv.gz")
        # mtime=0 prevents volatile gzip headers, preserving reproducible extracted hashes.
        import io
        payload=io.StringIO(newline=""); writer=csv.DictWriter(payload,fieldnames=["excel_row","x","y"])
        writer.writeheader(); writer.writerows(rows)
        with target.open("wb") as stream:
            with gzip.GzipFile(fileobj=stream,mode="wb",mtime=0,filename="") as handle:
                handle.write(payload.getvalue().encode("utf-8"))
        manifests.append({"trace_id":name,"source_filename":filename,"sheet":sheetname,"excel_columns":[colx+1,coly+1],"n_pairs":len(rows),"source_sha256":hashlib.sha256(path.read_bytes()).hexdigest(),"extracted_sha256":hashlib.sha256(target.read_bytes()).hexdigest(),"extraction_rule":"Every numeric x/y pair; no filtering or interpolation; source row retained"})
        book.close()
    dump(destination/"extraction_manifest.json",manifests)
    # Preserve ambiguous stored means and untransformed numerical columns distinctly.
    wb=openpyxl.load_workbook(source/mappings[0][0],data_only=True,read_only=True)
    source_rows=list(wb["Figure2c"].values)
    audit=[]
    for rownumber in (2,3,4,6,7,8):
        r=source_rows[rownumber-1]; numeric=list(r[1:5])
        audit.append(dict(source="Nature2023 Figure2c",excel_row=rownumber,electrode="DSA" if rownumber<5 else "NCOOH",current_density_kA_m2=r[0],B=numeric[0],C=numeric[1],D=numeric[2],E=numeric[3],untransformed_arithmetic_mean=float(np.mean(numeric)),stored_mean_F=r[5],stored_sd_G=r[6],stored_minus_raw_mean=float(r[5]-np.mean(numeric)),interpretation="Stored metric extrapolated to ODC; raw-column transformation unresolved; not a demonstrated paper error"))
    write_csv(destination/"nature2023_energy_source_audit.csv",audit); wb.close()


def analyze_extracted(root: Path,out: Path):
    source=root/"data/literature"
    metadata=json.loads((source/"trace_metadata.json").read_text(encoding="utf-8"))
    metrics=[]; binned=[]
    fig,axes=plt.subplots(1,3,figsize=(13,4))
    for ax,meta in zip(axes,metadata):
        filename=source/(meta["trace_id"]+".csv.gz")
        with gzip.open(filename,"rt",encoding="utf-8",newline="") as handle:
            rows=list(csv.DictReader(handle))
        x=np.array([float(r["x"]) for r in rows]); y=np.array([float(r["y"]) for r in rows])
        if np.any(np.diff(x)<0): raise ValueError("Non-monotonic source time: do not silently sort")
        bins=np.floor(x).astype(int); local=[]
        for hour in np.unique(bins):
            ys=y[bins==hour]
            local.append(dict(trace_id=meta["trace_id"],hour=int(hour),point_count=int(len(ys)),median=float(np.median(ys)),minimum=float(min(ys)),maximum=float(max(ys)),q025=float(np.quantile(ys,.025)),q975=float(np.quantile(ys,.975))))
        binned.extend(local)
        first_mask=(x>=x[0])&(x<x[0]+1); last_mask=x>x[-1]-1
        first,last=float(np.median(y[first_mask])),float(np.median(y[last_mask]))
        intervals=np.diff(x); positive=intervals[intervals>0]
        starts=np.r_[0,np.flatnonzero(intervals!=0)+1]
        ends=np.r_[starts[1:],len(x)]
        unique_x=x[starts]
        unique_y=np.array([np.median(y[a:b]) for a,b in zip(starts,ends)])
        first_unique=float(np.median(unique_y[unique_x<unique_x[0]+1]))
        last_unique=float(np.median(unique_y[unique_x>unique_x[-1]-1]))
        # Entire trace descriptive OLS on hour medians; no iid-row confidence interval.
        slope=float(np.polyfit([r["hour"]+.5 for r in local],[r["median"] for r in local],1)[0])
        metrics.append(dict(trace_id=meta["trace_id"],source_DOI=meta["doi"],figure=meta["figure"],n_source_rows=len(x),independent_electrode_count="unverified_from_source_workbook",time_start_h=float(x[0]),time_end_h=float(x[-1]),duration_h=float(x[-1]-x[0]),first_value=float(y[0]),last_value=float(y[-1]),first_hour_median=first,last_hour_median=last,median_change=last-first,relative_median_change_percent=100*(last-first)/first,descriptive_hour_bin_slope_per_h=slope,median_sampling_interval_h=float(np.median(positive)),maximum_sampling_gap_h=float(max(intervals)),duplicate_time_count=int(sum(intervals==0)),y_unit=meta["y_unit"],censoring="end_of_record; failure event not inferable",evidence="EXECUTED_CALCULATION_FROM_PUBLISHED_NUMERICAL_DATA"))
        metrics[-1].update(unique_timestamps=len(unique_x),largest_equal_time_group=int(max(ends-starts)),first_hour_timestamp_balanced_median=first_unique,last_hour_timestamp_balanced_median=last_unique,timestamp_balanced_change=last_unique-first_unique,timestamp_balanced_relative_change_percent=100*(last_unique-first_unique)/first_unique)
        bx=np.array([r["hour"]+.5 for r in local]); median=np.array([r["median"] for r in local])
        ax.plot(bx,median,lw=1.5)
        ax.fill_between(bx,[r["q025"] for r in local],[r["q975"] for r in local],alpha=.25)
        ax.set(xlabel="Published time / h",ylabel=meta["y_label"],title=meta["short_title"])
        ax.text(.02,.02,"Band: within-hour 2.5–97.5%\nNot electrode uncertainty",transform=ax.transAxes,fontsize=7)
    fig.suptitle("Published source-data reanalysis | no lifetime extrapolation",fontsize=12)
    fig.tight_layout(); fig.savefig(out/"figures/published_stability.png",dpi=180); plt.close(fig)
    write_csv(out/"published_trace_metrics.csv",metrics); write_csv(out/"published_trace_hour_bins.csv",binned)
    dump(out/"published_trace_summary.json",{"evidence":"Published numerical data, locally recomputed descriptive statistics", "traces":metrics,"inference_limit":"Time points are not independent electrodes; no mechanistic fit, hazard model, electrode population CI or industrial lifetime claim"})
    return {"trace_count":len(metrics),"source_rows":sum(r["n_source_rows"] for r in metrics)}


if __name__=="__main__":
    import argparse
    p=argparse.ArgumentParser(description=__doc__); p.add_argument("--source",required=True,type=Path); p.add_argument("--destination",required=True,type=Path)
    a=p.parse_args(); extract_workbooks(a.source,a.destination)
