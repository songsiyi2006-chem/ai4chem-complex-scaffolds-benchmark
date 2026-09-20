"""New published-data calculations; original frozen Phase28 outputs stay unchanged."""
from __future__ import annotations
import argparse,csv,gzip,hashlib,json,os
from pathlib import Path
for key in ("OMP_NUM_THREADS","MKL_NUM_THREADS","OPENBLAS_NUM_THREADS"): os.environ[key]="2"
import numpy as np
from scipy.stats import t
from scipy.interpolate import PchipInterpolator
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

ROOT=Path(__file__).resolve().parents[1]

def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()
def dump(path,obj):path.write_text(json.dumps(obj,ensure_ascii=False,indent=2,allow_nan=False)+"\n",encoding="utf-8")
def write_csv(path,rows):
    rows=list(rows)
    with path.open("w",encoding="utf-8",newline="") as f:
        w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)
def read_csv(path):
    with path.open(encoding="utf-8",newline="") as f:return list(csv.DictReader(f))

def mean_interval(mean,sd,n=4,level=.95):
    if n<2 or sd<0:raise ValueError("Need n>=2 and nonnegative SD")
    half=float(t.ppf((1+level)/2,n-1)*sd/np.sqrt(n))
    return mean-half,mean+half

def contrast(m0,s0,m1,s1,n=4,rho=None,level=.95):
    """Summary-statistic interval. Pairing/correlation assumptions remain explicit."""
    delta=m1-m0
    if rho is None:
        v0,v1=s0*s0/n,s1*s1/n
        variance=v0+v1
        df=variance**2/(v0*v0/(n-1)+v1*v1/(n-1)) if variance else float(n-1)
        label="independent_groups_assumption_unverified"
    else:
        if not -1<=rho<=1:raise ValueError("Correlation outside [-1,1]")
        variance=max(0,(s0*s0+s1*s1-2*rho*s0*s1)/n)
        df=float(n-1);label="paired_same4_assumption_unknown_rho"
    half=float(t.ppf((1+level)/2,df)*np.sqrt(variance))
    return {"difference":delta,"standard_error":float(np.sqrt(variance)),"df":float(df),"ci_low":delta-half,"ci_high":delta+half,"assumption":label}

def interpolate_current(potential,current,target):
    potential,current=np.asarray(potential),np.asarray(current)
    mask=current>0;potential,current=potential[mask],current[mask]
    if np.any(np.diff(current)<=0):raise ValueError("Current must be strictly increasing; no silent sorting")
    if target<current[0] or target>current[-1]:raise ValueError("Extrapolation is forbidden")
    idx=np.searchsorted(current,target);idx=max(1,min(idx,len(current)-1))
    linear=float(np.interp(target,current,potential))
    pchip=float(PchipInterpolator(current,potential)(target))
    return linear,pchip,float(current[idx]-current[idx-1])

def unique_time_medians(x,y):
    x,y=np.asarray(x),np.asarray(y)
    if np.any(np.diff(x)<0):raise ValueError("Source time order must be nondecreasing")
    starts=np.r_[0,np.flatnonzero(np.diff(x)!=0)+1];ends=np.r_[starts[1:],len(x)]
    return x[starts],np.array([np.median(y[a:b]) for a,b in zip(starts,ends)])

def window_change(x,y,exclude,width):
    start=float(x[0]+exclude);stop=float(x[-1])
    if start+width>stop-width:raise ValueError("Windows overlap or extend outside trace")
    a=y[(x>=start)&(x<start+width)];b=y[(x>stop-width)&(x<=stop)]
    if min(len(a),len(b))==0:raise ValueError("Empty window")
    av,bv=float(np.median(a)),float(np.median(b))
    return {"first_window_start_h":start,"first_window_end_h":start+width,"last_window_start_h":stop-width,"last_window_end_h":stop,"first_unique_times":len(a),"last_unique_times":len(b),"first_median":av,"last_median":bv,"change":bv-av,"relative_change_percent":100*(bv-av)/av}

def extract(workbook,destination):
    import openpyxl
    cfg=json.loads((ROOT/"configs/additional_public_analysis.json").read_text())
    if sha(workbook)!=cfg["source_workbook_sha256"]:raise ValueError("Unexpected source workbook hash")
    if destination.exists():raise ValueError("Use new extraction directory")
    destination.mkdir(parents=True)
    book=openpyxl.load_workbook(workbook,data_only=True,read_only=True)
    summaries=[]
    for sheet,quantity,unit in [("2e","Cl2_production","mol"),("2f","FE","percent")]:
        for row_index,row in enumerate(book[sheet].iter_rows(values_only=True),1):
            if isinstance(row[0],(int,float)):
                summaries.append(dict(sheet=sheet,excel_row=row_index,time_h=row[0],quantity=quantity,unit=unit,mean=row[1],sd=row[2],reported_n=4,raw_replicates="not_available",pairing="not_available"))
    write_csv(destination/"production_fe_summary_inputs.csv",summaries)
    curves=[]
    for sheet,colx,coly,label in [("2a",0,1,"before_crosspanel_iR_corrected"),("2a",3,4,"before_crosspanel_uncorrected"),("2h",3,4,"after100h_iR_unverified"),("2h",6,7,"after200h_iR_unverified")]:
        for row_index,row in enumerate(book[sheet].iter_rows(values_only=True),1):
            if isinstance(row[colx],(int,float)) and isinstance(row[coly],(int,float)):
                curves.append(dict(curve_id=label,sheet=sheet,excel_row=row_index,excel_potential_column=colx+1,excel_current_column=coly+1,potential_V_NHE=row[colx],current_mA_cm2=row[coly]))
    write_csv(destination/"polarization_inputs.csv",curves)
    dump(destination/"source_manifest.json",{"source_doi":"10.1038/s44160-026-01039-y","source_url":"https://media.springernature.com/original/springer-static/esm/art%3A10.1038%2Fs44160-026-01039-y/MediaObjects/44160_2026_1039_MOESM4_ESM.xlsx","source_sha256":sha(workbook),"figure_caption":"https://www.nature.com/articles/s44160-026-01039-y/figures/2","figure_review":"Published image visually inspected for mol Cl2, FE percent, h and V vs NHE/mA cm^-2 axes","n_definition":"Mean plus/minus SD of four independent experiments, Figure2e/f caption; SI printedp1/PDFp4 general statement","raw_replicates_available":False,"summary_rows":len(summaries),"polarization_points":len(curves),"extracted_files":[{"path":p.name,"sha256":sha(p)} for p in sorted(destination.glob('*.csv'))]})

def run(out):
    cfgpath=ROOT/"configs/additional_public_analysis.json";cfg=json.loads(cfgpath.read_text())
    inputs=ROOT/"data/additional_public";summary_inputs=read_csv(inputs/"production_fe_summary_inputs.csv")
    if cfg['reported_replicates']!=4 or any(int(r['reported_n'])!=cfg['reported_replicates'] for r in summary_inputs):
        raise ValueError("Published sample size is fixed at four; config and source rows must agree")
    if not 0<cfg['confidence_level']<1:raise ValueError("Confidence level must lie in (0,1)")
    out.mkdir(parents=True);figures=out/"figures";figures.mkdir()
    intervals=[];lookup={}
    for row in summary_inputs:
        mean,sd=float(row['mean']),float(row['sd']);hour=float(row['time_h']);n=int(row['reported_n'])
        lo,hi=mean_interval(mean,sd,n,cfg['confidence_level'])
        intervals.append(dict(quantity=row['quantity'],time_h=hour,unit=row['unit'],mean=mean,sd=sd,n=n,ci95_low=lo,ci95_high=hi,assumption="Normal independent replicate means; timepoints not independent",zero_origin="fixed_origin_not_stochastic" if hour==0 and mean==0 and sd==0 else "no"))
        lookup[row['quantity'],hour]=(mean,sd)
    write_csv(out/"replicate_mean_intervals.csv",intervals)
    contrasts=[]
    for quantity in ("Cl2_production","FE"):
        for a,b in cfg['contrasts_h']:
            m0,s0=lookup[quantity,a];m1,s1=lookup[quantity,b]
            for rho in [None]+cfg['unknown_within_run_correlations']:
                result=contrast(m0,s0,m1,s1,cfg['reported_replicates'],rho,cfg['confidence_level'])
                if quantity=='Cl2_production' and a==0:
                    # Origin is fixed zero: endpoint t interval, no manufactured paired df.
                    result['assumption']='fixed_zero_origin_endpoint_interval'
                contrasts.append(dict(quantity=quantity,time_start_h=a,time_end_h=b,rho="independent" if rho is None else rho,**result))
    write_csv(out/"time_contrasts_correlation_sensitivity.csv",contrasts)
    curve_inputs=read_csv(inputs/"polarization_inputs.csv");curves={}
    for row in curve_inputs:curves.setdefault(row['curve_id'],[]).append((float(row['potential_V_NHE']),float(row['current_mA_cm2'])))
    interpolated=[];values={}
    for label,pairs in curves.items():
        pairs=np.array(pairs)
        for current in cfg['current_grid_mA_cm2']:
            linear,pchip,bracket=interpolate_current(pairs[:,0],pairs[:,1],current)
            interpolated.append(dict(curve_id=label,current_mA_cm2=current,linear_potential_V_NHE=linear,pchip_potential_V_NHE=pchip,method_difference_mV=1000*(pchip-linear),bracket_width_mA_cm2=bracket,calibration_status="crosspanel_iR_pairing_unverified" if label.startswith('before') else "sharedpanel_iR_unverified"))
            values[label,current]=linear
    write_csv(out/"polarization_interpolated.csv",interpolated)
    differences=[]
    for current in cfg['current_grid_mA_cm2']:
        after100,after200=values['after100h_iR_unverified',current],values['after200h_iR_unverified',current]
        for baseline in ('after100h_iR_unverified','before_crosspanel_iR_corrected','before_crosspanel_uncorrected'):
            differences.append(dict(baseline=baseline,comparison="after200h_iR_unverified",current_mA_cm2=current,coordinate_difference_mV=1000*(after200-values[baseline,current]),scientific_interpretation="descriptive_sharedpanel_coordinate_change_not_calibrated_kinetics" if baseline.startswith('after') else "crosspanel_sensitivity_only_zero_hour_pairing_not_established"))
    write_csv(out/"polarization_differences.csv",differences)
    robustness=[]
    for meta in json.loads((ROOT/'data/literature/trace_metadata.json').read_text()):
        with gzip.open(ROOT/'data/literature'/(meta['trace_id']+'.csv.gz'),'rt') as f:raw=list(csv.DictReader(f))
        x,y=unique_time_medians([float(r['x']) for r in raw],[float(r['y']) for r in raw])
        for exclusion in cfg['startup_exclusion_h']:
            for width in cfg['endpoint_window_h']:
                robustness.append(dict(trace_id=meta['trace_id'],startup_exclusion_h=exclusion,window_h=width,y_unit=meta['y_unit'],**window_change(x,y,exclusion,width)))
    write_csv(out/"stability_window_robustness.csv",robustness)
    # Descriptive production fit only: correlated cumulative-time points are not replicates.
    amount=[r for r in intervals if r['quantity']=='Cl2_production']
    hours=np.array([r['time_h'] for r in amount]);means=np.array([r['mean'] for r in amount]);slope=float(hours@means/(hours@hours))
    final=lookup['Cl2_production',24];final_ci=mean_interval(*final,n=cfg['reported_replicates'],level=cfg['confidence_level'])
    headline={'evidence':'EXECUTED_CALCULATION_FROM_PUBLISHED_SUMMARY_AND_TRACE_DATA','config_sha256':sha(cfgpath),'published_summary_rows':len(summary_inputs),'new_polarization_points':len(curve_inputs),'mean_intervals':len(intervals),'contrast_sensitivity_rows':len(contrasts),'polarization_grid_rows':len(interpolated),'stability_robustness_rows':len(robustness),'final_Cl2_mol':final[0],'final_Cl2_mean_ci95_mol':list(final_ci),'average_24h_rate_mol_h':final[0]/24,'average_24h_rate_mean_ci95_mol_h':[v/24 for v in final_ci],'descriptive_cumulative_rate_origin_OLS_mol_h':slope,'descriptive_origin_fit_max_residual_mol':float(max(abs(means-hours*slope))),'final_FE_percent':lookup['FE',24][0],'final_FE_mean_ci95_percent':list(mean_interval(*lookup['FE',24],n=cfg['reported_replicates'],level=cfg['confidence_level'])),'FE_change_24minus0_percentage_points':lookup['FE',24][0]-lookup['FE',0][0],'FE_change_24minus1_percentage_points':lookup['FE',24][0]-lookup['FE',1][0],'after200minus100_at1000mA_mV':next(r['coordinate_difference_mV'] for r in differences if r['baseline'].startswith('after') and r['current_mA_cm2']==1000),'maximum_linear_PCHIP_difference_mV':max(abs(r['method_difference_mV']) for r in interpolated),'blind_or_experiment_status':'No new experiments, no replicate-level reconstruction, no mechanism identification'}
    headline['stability_ranges']=[]
    for trace in sorted({r['trace_id'] for r in robustness}):
        subset=[r for r in robustness if r['trace_id']==trace]
        late=[r for r in subset if r['startup_exclusion_h']>=10]
        headline['stability_ranges'].append({'trace_id':trace,'all_grid_change_min_max':[min(r['change'] for r in subset),max(r['change'] for r in subset)],'startup_at_least10h_change_min_max':[min(r['change'] for r in late),max(r['change'] for r in late)],'all_grid_relative_percent_min_max':[min(r['relative_change_percent'] for r in subset),max(r['relative_change_percent'] for r in subset)]})
    dump(out/'summary.json',headline)
    # Three plots segregate replicate-based intervals from within-trace sensitivity.
    fig,axes=plt.subplots(1,3,figsize=(14,4))
    for ax,q in zip(axes[:2],('Cl2_production','FE')):
        subset=[r for r in intervals if r['quantity']==q]
        ax.errorbar([r['time_h'] for r in subset],[r['mean'] for r in subset],yerr=[r['mean']-r['ci95_low'] for r in subset],fmt='o-',ms=3,capsize=2)
        ax.set(xlabel='Time / h',ylabel='Cl2 / mol' if q=='Cl2_production' else 'FE / %',title='Pointwise t intervals, n=4 (not SD bars)')
    for rho in cfg['unknown_within_run_correlations']:
        z=next(r for r in contrasts if r['quantity']=='FE' and r['time_start_h']==0 and r['rho']==rho)
        axes[2].errorbar(rho,z['difference'],yerr=[[z['difference']-z['ci_low']],[z['ci_high']-z['difference']]],fmt='o',capsize=3)
    axes[2].axhline(0,color='gray',linestyle='--');axes[2].set(xlabel='Unknown paired correlation (assumed)',ylabel='FE change 24h - 0h / percentage points',title='Conclusion depends on missing covariance')
    fig.tight_layout();fig.savefig(figures/'production_FE_uncertainty.png',dpi=180);plt.close(fig)
    fig,axes=plt.subplots(1,2,figsize=(11,4))
    for label,pairs in curves.items():
        a=np.array(pairs);mask=a[:,1]>=0;axes[0].plot(a[mask,0],a[mask,1],label=label.replace('_',' '))
    axes[0].set(xlabel='Anode potential / V vs NHE',ylabel='Current / mA cm$^{-2}$',title='Source curves; cross-panel baseline unmatched');axes[0].legend(fontsize=6)
    for baseline in ('after100h_iR_unverified','before_crosspanel_iR_corrected','before_crosspanel_uncorrected'):
        z=[r for r in differences if r['baseline']==baseline];axes[1].plot([r['current_mA_cm2'] for r in z],[r['coordinate_difference_mV'] for r in z],'o-',label=baseline.replace('_',' '))
    axes[1].axhline(0,color='gray',linestyle='--');axes[1].set(xlabel='Common current / mA cm$^{-2}$',ylabel='After200h minus baseline / mV',title='Coordinate differences; no cell-energy inference');axes[1].legend(fontsize=6)
    fig.tight_layout();fig.savefig(figures/'polarization_matched_currents.png',dpi=180);plt.close(fig)
    fig,axes=plt.subplots(1,3,figsize=(13,4))
    for ax,trace in zip(axes,sorted({r['trace_id'] for r in robustness})):
        is_current='potential' in trace
        for width in cfg['endpoint_window_h']:
            z=[r for r in robustness if r['trace_id']==trace and r['window_h']==width]
            ax.plot([r['startup_exclusion_h'] for r in z],[r['relative_change_percent'] if is_current else 1000*r['change'] for r in z],'o-',label=str(width)+'h windows')
        ax.axhline(0,color='gray',linestyle='--');ax.set(xlabel='Startup excluded / h after first source time',ylabel='Relative ordinate change / %' if is_current else 'Anode potential change / mV',title=trace.replace('_',' '));ax.legend(fontsize=7)
    fig.suptitle('Unique-timestamp sensitivity grid; no electrode-population intervals',fontsize=11)
    fig.tight_layout();fig.savefig(figures/'startup_window_sensitivity.png',dpi=180);plt.close(fig)
    dump(out/'input_output_manifest.json',{'config_sha256':sha(cfgpath),'inputs':[{'path':str(p.relative_to(ROOT)).replace('\\','/'),'sha256':sha(p)} for p in sorted(inputs.glob('*')) if p.is_file() and p.suffix in ('.csv','.json')]+[{'path':str(p.relative_to(ROOT)).replace('\\','/'),'sha256':sha(p)} for p in sorted((ROOT/'data/literature').glob('*.csv.gz'))],'outputs':[{'path':str(p.relative_to(out)).replace('\\','/'),'sha256':sha(p)} for p in sorted(out.rglob('*')) if p.is_file()]})
    return headline

def main():
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--out',type=Path);p.add_argument('--extract-workbook',type=Path);p.add_argument('--extract-to',type=Path)
    a=p.parse_args()
    if a.extract_workbook:
        if a.extract_to is None:p.error('--extract-to required')
        extract(a.extract_workbook,a.extract_to);return
    if a.out is None or a.out.exists():p.error('--out must be a new directory')
    print(json.dumps(run(a.out.resolve()),ensure_ascii=False))

if __name__=='__main__':main()
